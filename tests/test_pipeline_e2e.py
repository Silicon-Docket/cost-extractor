import zipfile
from decimal import Decimal
from pathlib import Path

import openpyxl

from cost_extractor.money_parser import build_custom_rule, default_rules
from cost_extractor.pipeline import run_pipeline
from cost_extractor.report import build_workbook, save_workbook


def test_full_pipeline_over_mixed_fixtures_default_rules(
    tmp_path: Path, simple_docx, text_pdf, corrupt_docx, password_protected_pdf
):
    # A zip bundling a couple of the same fixtures, to exercise ingestion's
    # zip path inside the full pipeline too.
    zip_path = tmp_path / "bundle.zip"
    with zipfile.ZipFile(zip_path, "w") as zf:
        zf.write(simple_docx, "zipped_invoice.docx")
        zf.writestr("notes.txt", "irrelevant")

    inputs = [simple_docx, text_pdf, corrupt_docx, password_protected_pdf, zip_path]
    result = run_pipeline(inputs, default_rules())

    # simple.docx ($1,234.56 + $500) + text.pdf ($1,234.56)
    # + zipped copy of simple.docx ($1,234.56 + $500)
    # corrupt.docx and the password-protected pdf contribute $0 (errors).
    expected_total = (
        Decimal("1734.56") + Decimal("1234.56") + Decimal("1734.56")
    )
    assert result.grand_total == expected_total

    out_path = tmp_path / "report.xlsx"
    save_workbook(build_workbook(result), out_path)
    reloaded = openpyxl.load_workbook(out_path, data_only=True)
    summary = reloaded["Summary"]
    grand_total_row = next(
        row
        for row in summary.iter_rows(values_only=True)
        if row[0] == "Grand Total"
    )
    assert Decimal(str(grand_total_row[3])) == expected_total


def test_full_pipeline_with_a_custom_rule_enabled(tmp_path: Path, simple_docx):
    docx_path = tmp_path / "euro_invoice.docx"
    # Reuse simple_docx's generation logic isn't available here directly,
    # so build a tiny docx inline via python-docx for a EUR-only amount.
    from docx import Document

    doc = Document()
    doc.add_paragraph("Cost: 45.00 EUR total")
    doc.save(docx_path)

    rules = default_rules()
    rules.append(build_custom_rule(r"(?P<amount>\d+(?:\.\d{2})?)\s?EUR", "Euro", 0))

    result = run_pipeline([docx_path], rules)

    assert result.grand_total == Decimal("45.00")
    assert result.documents[0].matches[0].rule_id == "custom_0"
